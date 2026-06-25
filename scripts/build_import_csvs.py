r"""
build_import_csvs.py
--------------------
Transforms the raw Test_Data_.xlsx (three sheets with three different layouts)
into ONE unified set of CSV files that match the public.test_records schema.

Unified columns (order matters for the psql \copy command):
    source, record_date, product_model, station, result,
    bursts, power_dbm, burst_amps, standby_amps, details

- result is the TEXT 'Pass' / 'Fail'  (1 = Pass, 0 = Fail in the source)
- details is a JSON object holding the sheet-specific extra columns
- Empty values are written as truly empty fields so Postgres COPY reads them as NULL
- Hand_Held rows with no timestamp are junk/padding and are skipped
"""

import openpyxl
import csv
import json
import os
import datetime

SRC = r"C:\Users\kmaya\Downloads\Test_Data_.xlsx"
OUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "import_files")
os.makedirs(OUT_DIR, exist_ok=True)

# Final column order written to every CSV
COLUMNS = ["source", "record_date", "product_model", "station", "result",
           "bursts", "power_dbm", "burst_amps", "standby_amps", "details"]


# Values that look like data but really mean "no value"
_NULLISH = {"", "null", "na", "n/a", "#n/a", "nan", "none"}


def is_blank(v):
    """True for None, empty, or text placeholders like 'NULL'/'NA'."""
    return v is None or (isinstance(v, str) and v.strip().lower() in _NULLISH)


def clean_str(v):
    """Trim strings; turn blanks/placeholders into None."""
    if is_blank(v):
        return None
    return str(v).strip()


def as_int(v):
    """Coerce to int for integer columns; blanks/placeholders -> None."""
    if is_blank(v):
        return None
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return None


def fmt_date(v):
    """Format a datetime as an ISO timestamp Postgres can read; None if missing."""
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if is_blank(v):
        return None
    # occasionally a date may arrive as text already
    return str(v).strip() or None


def pf_to_text(v):
    """1 -> 'Pass', 0 -> 'Fail'. Anything else -> None (row will be skipped)."""
    if v == 1:
        return "Pass"
    if v == 0:
        return "Fail"
    return None


def num(v):
    """Pass real numbers through (double precision columns); blanks become None."""
    if is_blank(v):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def details_json(d):
    """Build a compact JSON object, dropping None/placeholder values."""
    clean = {k: (v.strftime("%Y-%m-%d %H:%M:%S") if isinstance(v, datetime.datetime) else v)
             for k, v in d.items() if not is_blank(v)}
    return json.dumps(clean, separators=(",", ":"))


def map_hand_held(r):
    # TimeDate, TestStationID, Fixture, Trial, Pass/Fail, StartN, EndN,
    # Version, CSum, 406PWR, BAmps, Bursts, _121PF, STBAmps
    (time_date, station_id, fixture, trial, pf, start_n, end_n,
     version, csum, pwr, bamps, bursts, pf121, stbamps) = r
    if time_date is None:        # skip junk/padding rows
        return None
    result = pf_to_text(pf)
    if result is None:
        return None
    return {
        "source": "Hand_Held",
        "record_date": fmt_date(time_date),
        "product_model": None,                       # Hand_Held has no model column
        "station": clean_str(station_id),
        "result": result,
        "bursts": as_int(bursts),
        "power_dbm": num(pwr),
        "burst_amps": num(bamps),
        "standby_amps": num(stbamps),
        "details": details_json({
            "fixture": fixture, "trial": trial, "start_n": start_n, "end_n": end_n,
            "version": clean_str(version), "csum": csum, "_121pf": pf121,
        }),
    }


def map_fly(r):
    # TestDateTime, Station, Bursts, Overall_PF, ID#1, ID#,
    # Standby_Amps, Burst_Amps, Power_Ouput_dBm
    (dt, station, bursts, pf, id1, idn, stby, bamps, pwr) = r
    if dt is None:
        return None
    result = pf_to_text(pf)
    if result is None:
        return None
    return {
        "source": "FLY",
        "record_date": fmt_date(dt),
        "product_model": None,                       # FLY model is embedded in station name
        "station": clean_str(station),
        "result": result,
        "bursts": as_int(bursts),
        "power_dbm": num(pwr),
        "burst_amps": num(bamps),
        "standby_amps": num(stby),
        "details": details_json({"id_1": id1, "id_n": idn}),
    }


def map_boat(r):
    # TestDateTime, Station, DUTModel, Bursts, Overall_PF, ID2, ID1
    (dt, station, model, bursts, pf, id2, id1) = r
    if dt is None:
        return None
    result = pf_to_text(pf)
    if result is None:
        return None
    return {
        "source": "Boat",
        "record_date": fmt_date(dt),
        "product_model": clean_str(model),
        "station": clean_str(station),
        "result": result,
        "bursts": as_int(bursts),
        "power_dbm": None,                           # Boat has no power/amps columns
        "burst_amps": None,
        "standby_amps": None,
        "details": details_json({"id2": id2, "id1": id1}),
    }


SHEETS = {
    "Hand_Held": ("import_Hand_Held.csv", map_hand_held),
    "FLY": ("import_FLY.csv", map_fly),
    "Boat": ("import_Boat.csv", map_boat),
}


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    grand_total = 0
    summary = []
    for sheet, (fname, mapper) in SHEETS.items():
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        next(rows)  # skip header
        out_path = os.path.join(OUT_DIR, fname)
        written = 0
        fails = 0
        with open(out_path, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(COLUMNS)
            for raw in rows:
                rec = mapper(raw)
                if rec is None:
                    continue
                w.writerow(["" if rec[c] is None else rec[c] for c in COLUMNS])
                written += 1
                if rec["result"] == "Fail":
                    fails += 1
        grand_total += written
        summary.append((sheet, fname, written, fails))
        print(f"{sheet:10s} -> {fname:24s}  rows={written:>7}  fails={fails:>6}")
    print(f"{'TOTAL':10s}    {'':24s}  rows={grand_total:>7}")
    return summary


if __name__ == "__main__":
    main()
